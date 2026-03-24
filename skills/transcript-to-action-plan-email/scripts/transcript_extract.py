#!/usr/bin/env python3
"""Transcript extractor for SBDC follow-up emails.

Goal: turn an uploaded transcript document (txt/docx/pdf) into a light-weight,
LLM-friendly JSON bundle. This script is intentionally *non-LLM* and focuses on:

- text extraction
- url + id detection
- naive sentence scoring for a quick "gist"
- heuristic candidate lists for next steps / accomplishments
- milestone keyword spotting (best-effort)

Usage:
  python scripts/transcript_extract.py --input path/to/transcript.docx

It prints JSON to stdout.
"""

from __future__ import annotations

import argparse
import json
import os
import re
from collections import Counter
from typing import Dict, List, Tuple


URL_RE = re.compile(r"https?://[^\s)\]}>\"']+", re.IGNORECASE)
CLIENT_ID_RE = re.compile(r"\b[A-Z]{2}\d{4}\b")


STOPWORDS = {
    "a",
    "an",
    "and",
    "are",
    "as",
    "at",
    "be",
    "but",
    "by",
    "for",
    "from",
    "has",
    "have",
    "he",
    "her",
    "hers",
    "him",
    "his",
    "i",
    "if",
    "in",
    "into",
    "is",
    "it",
    "its",
    "me",
    "my",
    "of",
    "on",
    "or",
    "our",
    "she",
    "so",
    "that",
    "the",
    "their",
    "them",
    "then",
    "there",
    "these",
    "they",
    "this",
    "to",
    "too",
    "up",
    "us",
    "was",
    "we",
    "were",
    "what",
    "when",
    "where",
    "which",
    "who",
    "will",
    "with",
    "you",
    "your",
}


MILESTONE_KEYWORDS: Dict[str, List[str]] = {
    "8(a) certification obtained": ["8(a)", "8a"],
    "dbe certified": ["dbe"],
    "edwosb certification obtained": ["edwosb"],
    "local disadvantaged business certification": ["disadvantaged business"],
    "mbe certified": ["mbe"],
    "mdot certification": ["mdot"],
    "sdb self-certified": ["sdb"],
    "wbe certified": ["wbe"],
    "wosb certification obtained": ["wosb"],
    "wbe/wosb": ["women owned", "woman owned"],
    "trademark obtained": ["trademark"],
    "business established": ["launched", "opened", "incorporated", "formed"],
    "business expansion": ["expanded", "new location", "second location"],
    "bought business": ["bought", "acquired"],
    "sold the business": ["sold the business", "sold my business"],
    "strategic growth plan success": ["growth plan"],
    "ai tools implemented": ["implemented ai", "chatgpt", "copilot", "gemini"],
    "ai improvement realized": ["time savings", "cost savings", "revenue", "customer satisfaction"],
}


def read_text(path: str) -> Tuple[str, str]:
    ext = os.path.splitext(path)[1].lower()
    if ext in {".txt", ".md"}:
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read(), "text"

    if ext == ".docx":
        try:
            import docx
        except Exception as e:
            raise RuntimeError("python-docx is required to read .docx") from e

        d = docx.Document(path)
        parts = []
        for p in d.paragraphs:
            if p.text:
                parts.append(p.text)
        for t in d.tables:
            for row in t.rows:
                for cell in row.cells:
                    txt = cell.text.strip()
                    if txt:
                        parts.append(txt)
        return "\n".join(parts), "docx"

    if ext == ".pdf":
        try:
            import pdfplumber
            text_parts = []
            with pdfplumber.open(path) as pdf:
                for page in pdf.pages:
                    t = page.extract_text() or ""
                    if t:
                        text_parts.append(t)
            return "\n".join(text_parts), "pdf"
        except Exception:
            try:
                from PyPDF2 import PdfReader
                reader = PdfReader(path)
                text_parts = []
                for page in reader.pages:
                    t = page.extract_text() or ""
                    if t:
                        text_parts.append(t)
                return "\n".join(text_parts), "pdf"
            except Exception as e:
                raise RuntimeError("Unable to read pdf; install pdfplumber or PyPDF2") from e

    raise ValueError(f"Unsupported file type: {ext} (supported: .txt, .docx, .pdf)")


def normalize(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[\t\u00A0]+", " ", text)
    text = re.sub(r" +", " ", text)
    return text.strip()


def split_sentences(text: str) -> List[str]:
    raw = re.split(r"(?<=[.!?])\s+", text)
    return [s.strip() for s in raw if len(s.strip()) >= 20][:5000]


def tokenize(text: str) -> List[str]:
    tokens = re.findall(r"[A-Za-z][A-Za-z']+", text.lower())
    return [t for t in tokens if t not in STOPWORDS and len(t) > 2]


def top_keywords(tokens: List[str], k: int = 20) -> List[Tuple[str, int]]:
    return Counter(tokens).most_common(k)


def score_sentences(sentences: List[str], freq: Counter) -> List[Tuple[str, float]]:
    scored = []
    for s in sentences:
        toks = tokenize(s)
        if not toks:
            continue
        score = sum(freq.get(t, 0) for t in toks) / (len(toks) ** 0.5)
        scored.append((s, score))
    scored.sort(key=lambda x: x[1], reverse=True)
    return scored


def heuristic_candidates(sentences: List[str]) -> Dict[str, List[str]]:
    ns_triggers = ("next", "action", "you should", "we should", "we will", "i will", "plan to", "need to", "follow up", "send", "schedule", "draft", "review")
    acc_triggers = ("completed", "done", "finished", "achieved", "obtained", "certified", "approved", "launched", "filed", "submitted", "signed")

    next_steps = []
    accomplishments = []
    for s in sentences:
        sl = s.lower()
        if any(t in sl for t in ns_triggers):
            next_steps.append(s)
        if any(t in sl for t in acc_triggers):
            accomplishments.append(s)

    return {
        "candidate_next_steps": next_steps[:25],
        "candidate_accomplishments": accomplishments[:25],
    }


def detect_milestones(text: str) -> List[str]:
    tl = text.lower()
    hits = []
    for milestone, kws in MILESTONE_KEYWORDS.items():
        for kw in kws:
            if kw.lower() in tl:
                hits.append(milestone)
                break
    return sorted(set(hits))


def detect_speakers(text: str) -> Dict[str, int]:
    counts: Counter = Counter()
    for line in text.split("\n"):
        m = re.match(r"^([A-Za-z][A-Za-z0-9 _\-]{0,40}):\s+", line.strip())
        if m:
            counts[m.group(1).strip()] += 1
    return dict(counts.most_common(10))


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", required=True, help="path to transcript (.txt, .docx, .pdf)")
    ap.add_argument("--top_sentences", type=int, default=8, help="number of gist sentences")
    args = ap.parse_args()

    text, ftype = read_text(args.input)
    text = normalize(text)

    urls = sorted(set(URL_RE.findall(text)))
    client_ids = sorted(set(CLIENT_ID_RE.findall(text)))

    sentences = split_sentences(text)
    tokens = tokenize(text)
    freq = Counter(tokens)
    scored = score_sentences(sentences, freq)
    gist = [s for s, _ in scored[: max(1, args.top_sentences)]]

    out = {
        "file_type": ftype,
        "char_count": len(text),
        "word_count": len(text.split()),
        "speaker_counts": detect_speakers(text),
        "client_id_candidates": client_ids,
        "urls": urls,
        "top_keywords": top_keywords(tokens, 20),
        "gist_sentences": gist,
        **heuristic_candidates(sentences),
        "milestone_keyword_hits": detect_milestones(text),
    }
    print(json.dumps(out, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
